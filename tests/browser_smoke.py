"""Real Chromium smoke test against an isolated local TeachBuddy instance."""
from __future__ import annotations
import io
import json
import os
from pathlib import Path
import secrets
import socket
import subprocess
import sys
import tempfile
import time
import urllib.request
import zipfile

from docx import Document
from pptx import Presentation
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts"
ARTIFACTS.mkdir(exist_ok=True)


def main():
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    base = f"http://127.0.0.1:{port}"
    errors, checks = [], []
    password = secrets.token_urlsafe(24)
    with tempfile.TemporaryDirectory(prefix="browser-data-", dir=ARTIFACTS) as data_dir:
        env = os.environ.copy()
        for key in list(env):
            if key.startswith("AI_") or key.startswith("TEACHBUDDY_"):
                env.pop(key)
        env.update(TEACHBUDDY_ENV="production", TEACHBUDDY_PASSWORD=password,
                   TEACHBUDDY_SECRET=secrets.token_urlsafe(48), TEACHBUDDY_SECURE_COOKIE="false",
                   TEACHBUDDY_DATA_DIR=data_dir, PYTHONDONTWRITEBYTECODE="1")
        with (ARTIFACTS / "browser-server.log").open("w", encoding="utf-8") as log:
            server = subprocess.Popen(
                [sys.executable, "-m", "uvicorn", "web.server:app", "--host", "127.0.0.1", "--port", str(port)],
                cwd=ROOT, env=env, stdout=log, stderr=log,
                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
            try:
                for _ in range(100):
                    try:
                        with urllib.request.urlopen(base + "/healthz", timeout=1):
                            break
                    except OSError:
                        if server.poll() is not None:
                            raise RuntimeError("Local test server exited; inspect browser-server.log")
                        time.sleep(0.1)
                else:
                    raise RuntimeError("Local test server did not become ready")
                with sync_playwright() as runner:
                    browser = runner.chromium.launch(channel="chrome", headless=True)
                    context = browser.new_context(viewport={"width": 1440, "height": 1100}, device_scale_factor=1, accept_downloads=True)
                    page = context.new_page()
                    page.on("pageerror", lambda error: errors.append(str(error)))
                    page.on("console", lambda message: errors.append(message.text) if message.type == "error" else None)
                    page.on("dialog", lambda dialog: dialog.accept())
                    page.goto(base)
                    expect(page.locator("#login-dialog")).to_be_visible()
                    page.locator("#login-password").fill(password)
                    page.locator("#login-button").click()
                    expect(page.locator("#login-dialog")).not_to_be_visible()
                    expect(page.locator("#lesson-template option")).to_have_count(3)
                    expect(page.locator("#ai-mode")).to_be_disabled()
                    checks.append("password login and offline status")
                    page.screenshot(path=str(ARTIFACTS / "web-desktop.png"), full_page=True)
                    page.locator("#lesson-title").fill("认识三角形")
                    page.locator("#lesson-subject").fill("数学")
                    page.locator("#lesson-grade").fill("四年级")
                    page.locator("#generate-button").click()
                    expect(page.locator("#result-content")).to_be_visible()
                    assert "认识三角形" in page.locator("#result-body").input_value()
                    checks.append("offline lesson generation")
                    body = page.locator("#result-body").input_value() + "\n课堂反馈：观察三角形。"
                    page.locator("#result-body").fill(body)
                    with page.expect_response(lambda response: response.url.endswith("/api/history") and response.request.method == "POST"):
                        page.locator("#save-result").click()
                    expect(page.locator("#result-status")).to_have_text("已保存")
                    checks.append("edit and save lesson")
                    for format_name in ("docx", "pptx", "txt"):
                        page.locator("#export-format").select_option(format_name)
                        with page.expect_download() as event:
                            page.locator("#export-result").click()
                        output = ARTIFACTS / ("browser-export." + format_name)
                        event.value.save_as(output)
                        if format_name == "docx":
                            document = Document(output)
                            assert any("认识三角形" in p.text for p in document.paragraphs)
                        elif format_name == "pptx":
                            assert len(Presentation(output).slides) >= 2
                        else:
                            assert "认识三角形" in output.read_text(encoding="utf-8-sig")
                    checks.append("Word, PowerPoint and text downloads open")
                    page.screenshot(path=str(ARTIFACTS / "web-result.png"), full_page=True)
                    page.locator('[data-nav="knowledge"]').click()
                    page.locator("#knowledge-file").set_input_files(
                        {"name": "triangle-notes.txt", "mimeType": "text/plain", "buffer": "三角形有三条边和三个角。".encode()})
                    expect(page.locator(".resource-row")).to_have_count(1)
                    page.locator(".resource-row input[type=checkbox]").check()
                    expect(page.locator("#knowledge-selected-count")).to_contain_text("1")
                    checks.append("knowledge upload and selection")
                    page.locator('[data-nav="files"]').click()
                    page.locator("#organize-files").set_input_files([
                        {"name": "lesson.txt", "mimeType": "text/plain", "buffer": b"same content"},
                        {"name": "copy.txt", "mimeType": "text/plain", "buffer": b"same content"},
                        {"name": "notes.md", "mimeType": "text/plain", "buffer": b"unique note"}
                    ])
                    page.locator("#preview-files").click()
                    expect(page.locator("#file-preview-list tr")).to_have_count(3)
                    page.locator("#exclude-duplicates").check()
                    with page.expect_download() as event:
                        page.locator("#download-files").click()
                    archive_path = ARTIFACTS / "browser-organized.zip"
                    event.value.save_as(archive_path)
                    with zipfile.ZipFile(archive_path) as archive:
                        assert len([name for name in archive.namelist() if not name.endswith("/")]) == 2
                    checks.append("file classification, duplicate preview and filtered ZIP")
                    page.locator('[data-nav="lesson"]').click()
                    page.reload()
                    expect(page.locator("#login-dialog")).not_to_be_visible()
                    expect(page.locator("#history-list .history-card")).to_have_count(2)
                    expect(page.locator("#lesson-template option")).to_have_count(3)
                    checks.append("authenticated refresh and persistent history")
                    page.set_viewport_size({"width": 390, "height": 844})
                    for section in ("lesson", "chat", "knowledge", "files"):
                        page.locator('[data-nav="' + section + '"]').click()
                        assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1"), section + " overflows on mobile"
                    page.locator('[data-nav="lesson"]').click()
                    page.screenshot(path=str(ARTIFACTS / "web-mobile.png"), full_page=True)
                    checks.append("four sections fit 390px mobile viewport")
                    page.locator("#open-settings").click()
                    expect(page.locator("#settings-dialog")).to_be_visible()
                    print('Starting logout navigation check', flush=True)
                    with page.expect_navigation(wait_until='domcontentloaded'):
                        page.locator("#logout-button").click()
                    expect(page.locator("#login-dialog")).to_be_visible()
                    checks.append("logout returns to protected login")
                    print('Logout verified; closing browser context', flush=True)
                    context.close()
                    print('Browser context closed; closing Chrome', flush=True)
                    browser.close()
                    print('Chrome closed', flush=True)
                assert not errors, errors
                report = {"passed": checks, "browser_errors": errors}
                (ARTIFACTS / "browser-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
                print(json.dumps(report, ensure_ascii=False))
            finally:
                server.terminate()
                try:
                    server.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    server.kill()
                    server.wait()


if __name__ == "__main__":
    main()