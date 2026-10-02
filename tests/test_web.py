"""Integration checks for the isolated browser API; no external AI calls."""
import importlib
import io
import json
import tempfile
import zipfile
from pathlib import PurePosixPath

import pytest
import httpx
import requests
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation

HEADERS = {"X-TeachBuddy-Request": "1"}
PASSWORD = "testing-password"
SECRET = "isolated-test-secret-must-have-at-least-32-characters"
TITLE = "\u5206\u6570\u7684\u521d\u6b65\u8ba4\u8bc6"
BODY = "\u4e00\u3001\u6559\u5b66\u76ee\u6807\nRecognize fractions.\n\u4e8c\u3001\u6559\u5b66\u8fc7\u7a0b\nFold a paper square."
REGULAR = "\u5e38\u89c4\u8bfe"
REVIEW = "\u590d\u4e60\u8bfe"
OPEN = "\u516c\u5f00\u8bfe"


@pytest.fixture
def make_app(monkeypatch, tmp_path):
    for key in ("AI_BASE_URL", "AI_API_KEY", "AI_MODEL", "AI_TIMEOUT",
                "AI_FALLBACK_BASE_URL", "AI_FALLBACK_API_KEY", "AI_FALLBACK_MODEL"):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("TEACHBUDDY_ENV", "development")
    monkeypatch.setenv("TEACHBUDDY_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("TEACHBUDDY_PASSWORD", PASSWORD)
    monkeypatch.setenv("TEACHBUDDY_SECRET", SECRET)
    monkeypatch.setenv("TEACHBUDDY_SECURE_COOKIE", "false")
    monkeypatch.setenv("TMP", str(tmp_path))
    monkeypatch.setenv("TEMP", str(tmp_path))
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))

    def no_network(*args, **kwargs):
        pytest.fail("An integration test attempted an external AI request.")
    monkeypatch.setattr(requests.sessions.Session, "request", no_network)
    monkeypatch.setattr(httpx.AsyncClient, "send", no_network)

    server = importlib.import_module("web.server")
    return server.create_app


def login(client):
    status = client.get("/api/status")
    assert status.status_code == 200
    response = client.post("/api/login", json={"password": PASSWORD}, headers=HEADERS)
    assert response.status_code == 200, response.text
    return client


@pytest.fixture
def app(make_app):
    return make_app()


@pytest.fixture
def client(app):
    with TestClient(app) as browser:
        yield login(browser)


def generated(client, template_id=REGULAR, **changes):
    payload = {"template_id": template_id, "title": TITLE, "subject": "Math",
               "grade": "Grade 3", "period": "40 min", "textbook": "Test edition",
               "requirements": "Use paper folding.", "mode": "offline"}
    payload.update(changes)
    response = client.post("/api/lesson/generate", json=payload, headers=HEADERS)
    assert response.status_code == 200, response.text
    return response.json()


def word_bytes(text="Knowledge paragraph", table_text="Table evidence"):
    document = Document()
    document.add_paragraph(text)
    document.add_table(rows=1, cols=1).cell(0, 0).text = table_text
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def upload_knowledge(client, name="notes.txt", data=b"Selected lesson evidence"):
    response = client.post("/api/knowledge", files={"file": (name, data)}, headers=HEADERS)
    assert response.status_code == 200, response.text
    return response.json()


def upload_template(client, content=None, name="lesson.json"):
    if content is None:
        content = json.dumps({"name": "My lesson", "sections": [
            {"title": "Learning goals", "default": "Discuss and compare fractions."}]}).encode()
    response = client.post("/api/templates/import", files={"file": (name, content)}, headers=HEADERS)
    assert response.status_code == 200, response.text
    return response.json()


def preview(client, uploads, **fields):
    response = client.post("/api/files/preview",
                           files=[("files", (name, body)) for name, body in uploads],
                           data=fields, headers=HEADERS)
    assert response.status_code == 200, response.text
    return response.json()


def download(client, batch, exclude=False):
    response = client.post("/api/files/" + batch["batch_id"] + "/download",
                           json={"exclude_duplicates": exclude}, headers=HEADERS)
    assert response.status_code == 200, response.text
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        assert archive.testzip() is None
        return {entry.filename: archive.read(entry) for entry in archive.infolist()}


def test_authentication_and_request_header(app):
    with TestClient(app) as browser:
        assert browser.get("/api/history").status_code == 401
        assert browser.post("/api/login", json={"password": PASSWORD}).status_code == 403
        assert browser.post("/api/login", json={"password": "wrong"}, headers=HEADERS).status_code == 401
        login(browser)
        assert browser.get("/api/history").status_code == 200
        assert browser.post("/api/history", json={"title": TITLE, "body": BODY}).status_code == 403
        assert browser.post("/api/history", json={"title": TITLE, "body": BODY},
                            headers={"X-TeachBuddy-Request": "0"}).status_code == 403
        response = browser.get("/api/status")
        assert response.headers["cache-control"] == "no-store"
        assert response.headers["x-content-type-options"] == "nosniff"
        assert "script-src 'self'" in response.headers["content-security-policy"]
        assert PASSWORD not in response.text
        assert SECRET not in response.text
        signed = browser.post("/api/login", json={"password": PASSWORD}, headers=HEADERS)
        assert "httponly" in signed.headers["set-cookie"].lower()


@pytest.mark.parametrize("template_id,heading", [
    (REGULAR, "\u6559\u5b66\u76ee\u6807"),
    (REVIEW, "\u590d\u4e60\u76ee\u6807"),
    (OPEN, "\u8bbe\u8ba1\u610f\u56fe"),
])
def test_all_builtin_templates_generate_and_save(client, template_id, heading):
    lesson = generated(client, template_id)
    assert lesson["mode"] == "offline"
    assert TITLE in lesson["body"]
    assert "Math" in lesson["body"]
    assert heading in lesson["body"]
    assert lesson["id"] in {item["id"] for item in client.get("/api/history").json()["items"]}


def test_manual_history_save_delete_and_session_isolation(client, app):
    response = client.post("/api/history", json={"title": TITLE, "body": BODY}, headers=HEADERS)
    assert response.status_code == 200, response.text
    item = response.json()
    with TestClient(app) as other:
        login(other)
        assert other.get("/api/history").json()["items"] == []
        assert other.delete("/api/history/" + item["id"], headers=HEADERS).status_code == 404
    assert client.delete("/api/history/" + item["id"], headers=HEADERS).status_code == 200
    assert client.get("/api/history").json()["items"] == []
    assert client.delete("/api/history/" + item["id"], headers=HEADERS).status_code == 404


@pytest.mark.parametrize("extension", ["docx", "pptx", "txt"])
def test_exports_are_valid_and_retain_lesson_content(client, extension):
    response = client.post("/api/export", json={"title": TITLE, "body": BODY, "format": extension},
                           headers=HEADERS)
    assert response.status_code == 200, response.text
    assert "attachment" in response.headers["content-disposition"]
    if extension == "docx":
        document = Document(io.BytesIO(response.content))
        text = "\n".join(p.text for p in document.paragraphs)
    elif extension == "pptx":
        presentation = Presentation(io.BytesIO(response.content))
        assert len(presentation.slides) >= 3
        text = "\n".join(shape.text for slide in presentation.slides
                         for shape in slide.shapes if shape.has_text_frame)
    else:
        text = response.content.decode("utf-8-sig")
    assert TITLE in text
    assert "Recognize fractions." in text
    assert "Fold a paper square." in text


@pytest.mark.parametrize("extension", ["txt", "md", "docx"])
def test_knowledge_import_delete_and_session_isolation(client, app, extension):
    data = word_bytes() if extension == "docx" else b"# Fraction evidence\n1/2 is one half."
    item = upload_knowledge(client, "evidence." + extension, data)
    assert item["chars"] > 10
    listed = client.get("/api/knowledge").json()["items"]
    assert item["id"] in {entry["id"] for entry in listed}
    assert all("text" not in entry for entry in listed)
    with TestClient(app) as other:
        login(other)
        assert other.get("/api/knowledge").json()["items"] == []
        assert other.delete("/api/knowledge/" + item["id"], headers=HEADERS).status_code == 404
        response = other.post("/api/lesson/generate", json={"template_id": REGULAR, "title": TITLE,
                              "knowledge_ids": [item["id"]]}, headers=HEADERS)
        assert response.status_code == 404
    assert client.delete("/api/knowledge/" + item["id"], headers=HEADERS).status_code == 200
    assert client.get("/api/knowledge").json()["items"] == []


@pytest.mark.parametrize("name,data", [
    ("empty.txt", b""),
    ("image.png", b"not text"),
    ("broken.docx", b"not a ZIP file"),
    ("control.txt", b"Lesson\x00secret"),
])
def test_bad_knowledge_does_not_create_records(client, name, data):
    response = client.post("/api/knowledge", files={"file": (name, data)}, headers=HEADERS)
    assert response.status_code in {400, 422}
    assert client.get("/api/knowledge").json()["items"] == []


def test_templates_import_generate_delete_and_session_isolation(client, app):
    builtin = client.get("/api/templates").json()["templates"]
    assert {REGULAR, REVIEW, OPEN}.issubset({item["id"] for item in builtin})
    item = upload_template(client)
    lesson = generated(client, item["id"])
    assert "Discuss and compare fractions." in lesson["body"]
    with TestClient(app) as other:
        login(other)
        assert item["id"] not in {entry["id"] for entry in other.get("/api/templates").json()["templates"]}
        response = other.post("/api/lesson/generate", json={"template_id": item["id"], "title": TITLE},
                              headers=HEADERS)
        assert response.status_code == 404
        assert other.delete("/api/templates/" + item["id"], headers=HEADERS).status_code == 404
    assert client.delete("/api/templates/" + item["id"], headers=HEADERS).status_code == 200


@pytest.mark.parametrize("content", [
    b"{bad JSON",
    b"[]",
    b'{"name":"Missing sections"}',
    b'{"name":"Empty sections","sections":[]}',
    b'{"name":"Wrong section","sections":[null]}',
    b'{"name":"Wrong type","sections":[{"title":123}]}',
])
def test_bad_template_json_returns_validation_error(client, content):
    before = client.get("/api/templates").json()
    response = client.post("/api/templates/import", files={"file": ("bad.json", content)}, headers=HEADERS)
    assert response.status_code in {400, 422}, response.text
    assert client.get("/api/templates").json() == before


@pytest.mark.parametrize("name", ["../../outside.json", "..\\outside.json", "/etc/passwd", "C:\\secret"])
def test_template_path_cannot_escape_builtin_store(client, name):
    response = client.post("/api/lesson/generate", json={"template_id": name, "title": TITLE}, headers=HEADERS)
    assert response.status_code in {400, 404, 422}


def test_file_preview_duplicate_removal_and_same_names_never_overwrite(client):
    uploads = [("lesson.txt", b"lesson-A"), ("lesson.txt", b"lesson-B"),
               ("copy.txt", b"lesson-A"), ("photo.png", b"photo-C")]
    batch = preview(client, uploads)
    assert len(batch["items"]) == 4
    assert batch["duplicate_count"] == 1
    all_files = download(client, batch)
    assert len(all_files) == 4
    assert sorted(all_files.values()) == sorted(body for _, body in uploads)
    unique_files = download(client, batch, exclude=True)
    assert len(unique_files) == 3
    assert set(unique_files.values()) == {b"lesson-A", b"lesson-B", b"photo-C"}
    for name in all_files:
        path = PurePosixPath(name)
        assert len(path.parts) == 2
        assert not path.is_absolute()
        assert ".." not in path.parts
        assert "\\" not in name


def test_batch_cannot_be_downloaded_from_other_session(client, app):
    batch = preview(client, [("notes.txt", b"private evidence")])
    with TestClient(app) as other:
        login(other)
        response = other.post("/api/files/" + batch["batch_id"] + "/download",
                              json={"exclude_duplicates": False}, headers=HEADERS)
        assert response.status_code == 404
    assert list(download(client, batch).values()) == [b"private evidence"]


@pytest.mark.parametrize("name", ["../../outside.txt", "..\\..\\outside.txt",
                                  "C:\\Windows\\secret.txt", "/etc/private.txt", "CON.txt"])
def test_uploaded_paths_are_safely_renamed_in_zip(client, name):
    response = client.post("/api/files/preview", files=[("files", (name, b"safe payload"))],
                           headers=HEADERS)
    if response.status_code in {400, 422}:
        return
    assert response.status_code == 200, response.text
    files = download(client, response.json())
    assert list(files.values()) == [b"safe payload"]
    for path in files:
        parts = PurePosixPath(path).parts
        assert len(parts) == 2
        assert not PurePosixPath(path).is_absolute()
        assert ".." not in parts
        assert "\\" not in path and ":" not in path
        assert parts[-1].split(".")[0].upper() != "CON"


@pytest.mark.parametrize("password,secret", [("", SECRET), (PASSWORD, ""), (PASSWORD, "short")])
def test_production_rejects_missing_or_weak_credentials(make_app, monkeypatch, password, secret):
    monkeypatch.setenv("TEACHBUDDY_ENV", "production")
    monkeypatch.setenv("TEACHBUDDY_PASSWORD", password)
    monkeypatch.setenv("TEACHBUDDY_SECRET", secret)
    with pytest.raises(RuntimeError):
        make_app()


def test_production_uses_secure_cookie_and_hides_docs(make_app, monkeypatch):
    monkeypatch.setenv("TEACHBUDDY_ENV", "production")
    monkeypatch.delenv("TEACHBUDDY_SECURE_COOKIE")
    application = make_app()
    with TestClient(application, base_url="https://testserver") as browser:
        response = browser.post("/api/login", json={"password": PASSWORD}, headers=HEADERS)
        assert response.status_code == 200
        assert "secure" in response.headers["set-cookie"].lower()
        assert response.headers["strict-transport-security"]
        assert browser.get("/docs").status_code == 404
        assert browser.get("/openapi.json").status_code == 404


def test_missing_ai_configuration_fails_cleanly(client):
    endpoints = [
        ("/api/lesson/generate", {"template_id": REGULAR, "title": TITLE, "mode": "ai"}),
        ("/api/lesson/revise", {"title": TITLE, "body": BODY, "instruction": "Add an activity"}),
        ("/api/chat", {"messages": [{"role": "user", "content": "Explain fractions"}]}),
    ]
    for path, payload in endpoints:
        response = client.post(path, json=payload, headers=HEADERS)
        assert response.status_code == 503, (path, response.text)


def test_ai_failure_hides_provider_details(make_app, monkeypatch):
    monkeypatch.setenv("AI_BASE_URL", "https://example.invalid/v1")
    monkeypatch.setenv("AI_MODEL", "test-model")
    server = importlib.import_module("web.server")

    def fail_provider(self, messages, **kwargs):
        raise RuntimeError("private-provider-key-private-provider-internal-trace")
    monkeypatch.setattr(server.AIClient, "chat", fail_provider)
    with TestClient(make_app()) as browser:
        login(browser)
        response = browser.post("/api/chat", json={"messages": [{"role": "user", "content": "Hello"}]},
                                headers=HEADERS)
        assert response.status_code == 502
        assert "private-provider" not in response.text


def test_custom_extension_and_filename_rules(client):
    batch = preview(client, [("normal.foo", b"custom extension"),
                             ("priority-note.txt", b"keyword priority")],
                    rules=json.dumps([{"category": "Custom", "exts": [".foo"]}]),
                    filename_rules=json.dumps([{"category": "Priority", "keywords": ["priority"]}]))
    files = download(client, batch)
    assert files["Custom/normal.foo"] == b"custom extension"
    assert files["Priority/priority-note.txt"] == b"keyword priority"


@pytest.mark.parametrize("rules", [
    "[invalid",
    '{"category":"not-an-array"}',
    '[{"category":"Custom","exts":["txt"]}]',
])
def test_invalid_file_rules_are_rejected(client, rules):
    response = client.post("/api/files/preview",
                           files=[("files", ("notes.txt", b"contents"))],
                           data={"rules": rules}, headers=HEADERS)
    assert response.status_code in {400, 422}, response.text


def test_ai_uses_only_selected_knowledge_and_preserves_docx_table_text(make_app, monkeypatch):
    monkeypatch.setenv("AI_BASE_URL", "https://example.invalid/v1")
    monkeypatch.setenv("AI_MODEL", "test-model")
    monkeypatch.setenv("AI_API_KEY", "private-test-api-key")
    server = importlib.import_module("web.server")
    calls = []

    def fake_chat(self, messages, **kwargs):
        calls.append(messages)
        return BODY
    monkeypatch.setattr(server.AIClient, "chat", fake_chat)
    with TestClient(make_app()) as browser:
        login(browser)
        selected = upload_knowledge(browser, "selected.docx",
                                    word_bytes("Selected paragraph evidence", "Selected table evidence"))
        upload_knowledge(browser, "unselected.txt", b"UNSELECTED PRIVATE CONTENT")
        lesson = generated(browser, mode="ai", knowledge_ids=[selected["id"]])
        assert "Recognize fractions." in lesson["body"]
        prompt = json.dumps(calls[-1], ensure_ascii=False)
        assert "Selected paragraph evidence" in prompt
        assert "Selected table evidence" in prompt
        assert "UNSELECTED PRIVATE CONTENT" not in prompt
        status = browser.get("/api/status").text
        assert "private-test-api-key" not in status
        response = browser.post("/api/chat", json={
            "messages": [{"role": "user", "content": "Explain this fraction"}],
            "knowledge_ids": [selected["id"]]}, headers=HEADERS)
        assert response.status_code == 200
        assert "Recognize fractions." in response.json()["reply"]
        assert "Selected table evidence" in json.dumps(calls[-1])


def test_ai_revision_returns_complete_lesson_and_saves_history(make_app, monkeypatch):
    monkeypatch.setenv("AI_BASE_URL", "https://example.invalid/v1")
    monkeypatch.setenv("AI_MODEL", "test-model")
    server = importlib.import_module("web.server")
    revised = "\u3010\u8ba8\u8bba\u8981\u70b9\u3011\nAdd collaborative work.\n\u3010\u66f4\u65b0\u6559\u6848\u3011\n" + BODY

    def fake_chat(self, messages, **kwargs):
        assert "Add collaborative work" in messages[-1]["content"]
        return revised
    monkeypatch.setattr(server.AIClient, "chat", fake_chat)
    with TestClient(make_app()) as browser:
        login(browser)
        response = browser.post("/api/lesson/revise", json={
            "title": TITLE, "body": BODY, "instruction": "Add collaborative work"}, headers=HEADERS)
        assert response.status_code == 200, response.text
        item = response.json()
        assert "Recognize fractions." in item["body"]
        assert item["mode"] == "ai"
        assert item["id"] in {entry["id"] for entry in browser.get("/api/history").json()["items"]}


def test_history_retains_the_twenty_newest_entries(client):
    identifiers = []
    for number in range(22):
        response = client.post("/api/history", json={
            "title": "Lesson " + str(number), "body": BODY}, headers=HEADERS)
        assert response.status_code == 200
        identifiers.append(response.json()["id"])
    saved = client.get("/api/history").json()["items"]
    assert [item["id"] for item in saved] == list(reversed(identifiers[-20:]))


def test_large_json_rejected_before_parsing(client):
    response = client.post("/api/history", content=b"x" * (1024 * 1024 + 1),
                           headers={**HEADERS, "Content-Type": "application/json"})
    assert response.status_code == 413


def test_logout_removes_access(client):
    generated(client)
    assert client.post("/api/logout", headers=HEADERS).status_code == 200
    assert client.get("/api/history").status_code == 401


def test_custom_category_paths_cannot_escape_archive_folder(client):
    batch = preview(client, [("notes.txt", b"safe custom category")],
                    rules=json.dumps([{"category": "../../escape", "exts": [".txt"]}]))
    files = download(client, batch)
    assert list(files.values()) == [b"safe custom category"]
    assert list(files) == ["escape/notes.txt"]
