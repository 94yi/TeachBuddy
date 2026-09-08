# -*- coding: utf-8 -*-
"""教案导出为 Word（python-docx），无依赖时退回纯文本。"""

import os


def _export_text(path: str, title: str, body: str) -> str:
    with open(path, "w", encoding="utf-8-sig") as handle:
        handle.write(title + "\n\n" + body)
    return path


def export_document(path: str, title: str, body: str) -> str:
    if os.path.splitext(path)[1].lower() != ".docx":
        return _export_text(os.path.splitext(path)[0] + ".txt", title, body)
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError:
        return _export_text(os.path.splitext(path)[0] + ".txt", title, body)

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.add_heading(title, level=0)
    for line in body.splitlines():
        text = line.strip()
        if not text:
            continue
        head_prefix = text[:2]
        numbered = head_prefix[:1] in "一二三四五六七八九十" and ("、" in text[:3] or "．" in text[:3])
        if numbered:
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    document.save(path)
    return path
