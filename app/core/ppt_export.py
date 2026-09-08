# -*- coding: utf-8 -*-
"""从教案正文生成基础演示文稿。"""

import os
import re


SECTION_RE = re.compile(r"^(?:#{1,3}\s*|[一二三四五六七八九十]+、\s*|（[一二三四五六七八九十]+）\s*)(.+)$")


def _sections(body: str):
    sections = []
    current = {"title": "", "lines": []}
    for raw in body.splitlines():
        text = raw.strip()
        if not text:
            continue
        match = SECTION_RE.match(text)
        if match:
            if current["title"] or current["lines"]:
                sections.append(current)
            current = {"title": match.group(1).strip("# \t"), "lines": []}
        else:
            current["lines"].append(text.lstrip("-• "))
    if current["title"] or current["lines"]:
        sections.append(current)
    if not sections and body.strip():
        sections.append({"title": "教学要点", "lines": body.strip().splitlines()})
    return sections


def build_outline(body: str) -> str:
    blocks = []
    for section in _sections(body):
        lines = ["# " + (section["title"] or "教学要点")]
        lines.extend("- " + line for line in section["lines"][:12])
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks)


def parse_outline(outline: str) -> list:
    slides = []
    for raw in outline.splitlines():
        text = raw.strip()
        if not text:
            continue
        if text.startswith("#"):
            title = text.lstrip("#").strip()
            slides.append({"title": title or "幻灯片", "lines": []})
            continue
        if not slides:
            slides.append({"title": "教学要点", "lines": []})
        point = text.lstrip("-• ").strip()
        if point:
            slides[-1]["lines"].append(point)
    if not slides:
        slides.append({"title": "教学要点", "lines": ["详见教案"]})
    return slides


def clean_outline_text(text: str) -> str:
    """清理 AI 返回的大纲：去掉代码块/强调/链接等标记，保留 # 标题与 - 要点结构。"""
    lines = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.strip()
        if line.startswith("```"):
            continue
        line = line.replace("**", "").replace("__", "").replace("`", "")
        line = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", line)
        lines.append(line)
    result = "\n".join(lines)
    result = re.sub(r"\n{3,}", "\n\n", result)
    return result.strip()


def extract_pptx_outline(path: str) -> str:
    """从已有 .pptx 提取每页标题与要点，生成可编辑的大纲文本。"""
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("读取 PPT 需要安装 python-pptx") from exc

    presentation = Presentation(path)
    blocks = []
    for slide in presentation.slides:
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            title = title_shape.text.strip()
        bullets = []
        for shape in slide.shapes:
            if shape == title_shape or not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    bullets.append(text)
        if not title and bullets:
            title = bullets.pop(0)
        if not title:
            continue
        block = ["# " + title]
        block.extend("- " + bullet for bullet in bullets[:12])
        blocks.append("\n".join(block))
    return "\n\n".join(blocks)


def export_outline(path: str, title: str, outline: str) -> str:
    if os.path.splitext(path)[1].lower() != ".pptx":
        path = os.path.splitext(path)[0] + ".pptx"
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("生成 PPT 需要安装 python-pptx") from exc

    presentation = Presentation()
    slide_layout = presentation.slide_layouts[1]
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_slide.placeholders[1].text = "教学演示文稿"

    for section in parse_outline(outline):
        if not section["title"] and not section["lines"]:
            continue
        slide = presentation.slides.add_slide(slide_layout)
        slide.shapes.title.text = section["title"] or title
        placeholder = slide.placeholders[1]
        text_frame = placeholder.text_frame
        text_frame.clear()
        first = True
        for raw_line in section["lines"][:12]:
            line = raw_line.strip()[:100]
            if not line:
                continue
            paragraph = text_frame.paragraphs[0] if first else text_frame.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(20)
            first = False
        if first:
            paragraph = text_frame.paragraphs[0]
            paragraph.text = "详见教案"
            paragraph.font.size = Pt(20)

    presentation.save(path)
    return path


def export_presentation(path: str, title: str, body: str) -> str:
    return export_outline(path, title, build_outline(body))
